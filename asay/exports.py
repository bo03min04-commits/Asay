"""증거 정리용 보고서: 공적 인증·전자증거 원본성 보장 서류가 아닙니다."""
from io import BytesIO
from xml.sax.saxutils import escape
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

FONT = "HYSMyeongJo-Medium"
try:
    pdfmetrics.getFont(FONT)
except KeyError:
    pdfmetrics.registerFont(UnicodeCIDFont(FONT))


def make_pdf(case, evidence) -> bytes:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=(210*mm,297*mm), topMargin=19*mm, bottomMargin=17*mm, leftMargin=18*mm, rightMargin=18*mm,
                            title=f"Asay Evidence - {case['title']}")
    body = ParagraphStyle('KrNormal', fontName=FONT, fontSize=9, leading=15, textColor=colors.HexColor('#222b3a'), spaceAfter=7)
    small = ParagraphStyle('KrSmall', parent=body, fontSize=7.6, leading=12)
    title = ParagraphStyle('KrTitle', parent=body, fontSize=18, leading=26, spaceAfter=15)
    h2 = ParagraphStyle('KrHeading', parent=body, fontSize=12, leading=19, spaceBefore=10, spaceAfter=9)
    story = [Paragraph('ASAY / 증거자료 정리본', title),
             Paragraph('본 자료는 사용자가 수집·입력한 정보를 정리한 참고문서이며, 증거능력·법적 판단·접수를 보장하지 않습니다.', small),
             Spacer(1, 8),
             Paragraph(f"사건: {escape(case['title'])}", h2),
             Paragraph(f"설명: {escape(case['description'] or '미입력')}", body),
             Paragraph(f"생성: {escape(case['created_at'])}  ·  증거 수: {len(evidence)}", small)]
    for i, item in enumerate(evidence, 1):
        details = [
            f"{i}. {escape(item['title'] or '제목 없음')}",
            f"URL: {escape(item['source_url'] or '수동 입력')}",
            f"플랫폼: {escape(item['platform'])} / 작성자: {escape(item['author'] or '미확인')}",
            f"게시 일시: {escape(item['posted_at'] or '미확인')} / 채증(기록) 일시: {escape(item['captured_at'])}",
            f"수집 상태: {escape(item['capture_status'])} / 레코드 SHA-256: {escape(item['sha256'])}",
            f"HTML SHA-256: {escape(item['html_sha256'])} / 첨부 SHA-256: {escape(item['attachment_sha256'])}",
        ]
        story.append(Paragraph(details[0], h2))
        story += [Paragraph(line, small) for line in details[1:]]
        story.append(Paragraph('내용 발췌 (최대 3000자)', small))
        excerpt = escape((item['body'] or '내용 없음')[:3000]).replace('\n', '<br/>')
        story.append(Paragraph(excerpt, body))
        if item['notes']:
            story.append(Paragraph('사용자 검토 메모: '+escape(item['notes'])[:1500], small))
        if item['warning']:
            story.append(Paragraph('주의: '+escape(item['warning'])[:500], small))
        story.append(Spacer(1, 8))
    def page_num(canvas, doc):
        canvas.setFont(FONT, 8)
        canvas.setFillColor(colors.HexColor('#738093'))
        canvas.drawString(18*mm, 10*mm, 'Asay · 사후 검토 및 원본 대조 필요')
        canvas.drawRightString(192*mm,10*mm,str(doc.page))
    doc.build(story, onFirstPage=page_num, onLaterPages=page_num)
    return out.getvalue()


def make_xlsx(case, evidence) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = '증거 목록'
    ws.append(['No.', '사건명', '플랫폼', '제목', '원본 URL', '작성자', '게시 일시', '수집 일시', '수집 상태', '내용', '검토 메모', 'SHA-256', '이전 채증 ID'])
    def excel_safe(value):
        if not isinstance(value, str):
            return value
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
        return value[:32000]
    for i, e in enumerate(evidence, 1):
        ws.append([excel_safe(v) for v in [i, case['title'], e['platform'], e['title'], e['source_url'], e['author'], e['posted_at'], e['captured_at'],
                   e['capture_status'], e['body'][:32000], e['notes'], e['sha256'], e['parent_id'] or '']])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [7,24,19,34,45,20,22,22,18,65,45,55,14]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17354A')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    note = wb.create_sheet('안내')
    note.append(['Asay 증거자료 관리 안내'])
    note.append(['자동 추출된 내용은 사이트 특성상 누락·오류가 있을 수 있으므로 원본과 비교하세요.'])
    note.append(['SHA-256은 저장된 데이터의 비교용 해시이며 원본성·증거능력 보장을 뜻하지 않습니다.'])
    note.append(['실제 신고 또는 법률상담 전에 인적사항·사실관계를 최종 검토하세요.'])
    note.column_dimensions['A'].width = 105
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def make_complaint_draft(case, evidence):
    lines = [
        '사실관계 정리 및 상담용 초안 (제출용 완성 고소장이 아닙니다)',
        '', f"사건명: {case['title']}", f"사건 설명: {case['description']}",
        '피해자 성명/연락처: [직접 입력]',
        '상대방(확인 가능한 정보): [직접 입력]',
        '구체적 피해 및 경위: [날짜, 상황, 피해를 사실 중심으로 입력]',
        '', '첨부 증거 목록',
    ]
    for i, e in enumerate(evidence, 1):
        lines.append(f"{i}. {e['title'] or '온라인 콘텐츠'} / {e['source_url'] or '수동 등록'} / 작성자 {e['author'] or '미확인'} / 게시 {e['posted_at'] or '미확인'}")
    lines += ['', '요청 사항: [관련 기관 또는 법률 전문가와 확인 후 작성]', '',
              '※ 죄명 또는 성립 여부는 자동으로 단정하지 않습니다. 제출 전 전문가 또는 담당 기관의 확인을 권장합니다.']
    return '\n'.join(lines)
